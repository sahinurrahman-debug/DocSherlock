import io
import json

import pytest

from app.services import ocr
from app.services.extractor import ExtractionError, extract, detect_doc_date

needs_ocr = pytest.mark.skipif(not ocr.available(), reason="OCR engine not installed")


def text_of(ex):
    return "\n".join(p.text for p in ex.pages)


def test_unsupported_and_empty():
    with pytest.raises(ExtractionError):
        extract("malware.exe", b"MZ...")
    with pytest.raises(ExtractionError):
        extract("empty.txt", b"")


def test_text_encodings():
    assert "caf\u00e9" in text_of(extract("a.txt", "caf\u00e9 ouvert".encode("utf-8")))
    assert "caf\u00e9" in text_of(extract("b.txt", "caf\u00e9 ouvert".encode("cp1252")))
    assert "hello" in text_of(extract("c.txt", "hello".encode("utf-16")))


def test_markdown_keeps_heading_markers():
    ex = extract("n.md", b"# Title\n\nBody text here.")
    assert text_of(ex).startswith("# Title")


def test_csv_header_units_survive():
    ex = extract("m.csv", b"Month,On-time rate (%),Cost ($)\nJanuary,96.2,1200\n")
    t = text_of(ex)
    assert "On-time rate: 96.2%" in t and "Cost: $1200" in t


def test_html_strips_scripts_and_keeps_headings():
    ex = extract("p.html", b"<html><title>T</title><h1>Head</h1><p>Body</p><script>evil()</script></html>")
    t = text_of(ex)
    assert "# Head" in t and "Body" in t and "evil" not in t
    assert ex.meta["title"] == "T"


def test_json_is_flattened():
    ex = extract("d.json", json.dumps({"a": {"b": [1, 2]}, "name": "x"}).encode())
    assert "a.b[1]: 2" in text_of(ex)


def test_invalid_json_falls_back_to_text():
    ex = extract("d.json", b"{not json")
    assert ex.warnings and "not json" in text_of(ex)


def test_eml_headers_date_and_attachment_warning():
    raw = (b"From: a@x.com\nTo: b@x.com\nDate: Fri, 02 Feb 2024 10:21:00 +0000\nSubject: Hi\nContent-Type: text/plain\n\nPayment is Net 30.\n")
    ex = extract("m.eml", raw)
    assert "Net 30" in text_of(ex)
    assert detect_doc_date("m.eml", ex) == ("2024-02-02", "email header")


def test_docx_headings_tables(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("Payment Terms", 1)
    d.add_paragraph("Invoices are payable Net 45.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Cost"
    t.cell(1, 0).text, t.cell(1, 1).text = "Widget", "$5"
    buf = io.BytesIO()
    d.save(buf)
    ex = extract("x.docx", buf.getvalue())
    txt = text_of(ex)
    assert "# Payment Terms" in txt and "Item: Widget; Cost: $5" in txt
    assert ex.paged is False


def test_xlsx_sheets(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Budget"
    wb.active.append(["Team", "Spend ($)"])
    wb.active.append(["Ops", 1200])
    buf = io.BytesIO()
    wb.save(buf)
    ex = extract("b.xlsx", buf.getvalue())
    assert "Team: Ops; Spend: $1200" in text_of(ex)


def make_pdf(pages: list[str]) -> bytes:
    import pymupdf
    doc = pymupdf.open()
    for t in pages:
        pg = doc.new_page()
        pg.insert_text((72, 100), "Heading One", fontsize=20)
        pg.insert_textbox(pymupdf.Rect(72, 130, 520, 400), t, fontsize=11)
    return doc.tobytes()


def test_pdf_text_pages_and_heading_detection():
    ex = extract("a.pdf", make_pdf(["Termination requires sixty (60) days written notice.", "Second page body text."]))
    assert len(ex.pages) == 2 and ex.paged
    assert "# Heading One" in ex.pages[0].text
    assert "sixty (60) days" in ex.pages[0].text


def test_pdf_corrupt_and_encrypted():
    with pytest.raises(ExtractionError):
        extract("bad.pdf", b"%PDF-1.4 this is not really a pdf")
    import pymupdf
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 100), "secret")
    enc = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u")
    with pytest.raises(ExtractionError, match="password"):
        extract("locked.pdf", enc)


def render_text_png(text: str, size=(1000, 220), font_size=40) -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", size, "white")
    ImageDraw.Draw(img).text((30, 60), text, fill="black", font=ImageFont.truetype("C:/Windows/Fonts/arial.ttf", font_size))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


@needs_ocr
def test_image_ocr_returns_text_and_confidence():
    ex = extract("scan.png", render_text_png("Headcount is 128 employees"))
    page = ex.pages[0]
    assert "128" in page.text and "employees" in page.text.lower()
    assert page.ocr_conf and page.ocr_conf > 0.8 and page.boxes


@needs_ocr
def test_blank_image_warns_instead_of_failing():
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (400, 300), "white").save(buf, "PNG")
    ex = extract("blank.png", buf.getvalue())
    assert not ex.pages[0].text.strip() and any("No text" in w for w in ex.warnings)


@needs_ocr
def test_scanned_pdf_page_uses_ocr():
    import pymupdf
    png = render_text_png("Liability is capped at $1,000,000", size=(1200, 300), font_size=44)
    doc = pymupdf.open()
    pg = doc.new_page(width=595, height=300)
    pg.insert_image(pg.rect, stream=png)
    ex = extract("scan.pdf", doc.tobytes())
    assert ex.pages[0].method == "ocr" and "1,000,000" in ex.pages[0].text


def test_doc_date_detection_priority():
    ex = extract("Agreement_2023-05-17.txt", b"Contract text. Dated: March 3, 2022.\nMore text.")
    assert detect_doc_date("Agreement_2023-05-17.txt", ex) == ("2022-03-03", "stated in document")
    ex = extract("Notes_2021-09-30.txt", b"Meeting notes without any date")
    assert detect_doc_date("Notes_2021-09-30.txt", ex) == ("2021-09-30", "file name")
    ex = extract("plain.txt", b"nothing here")
    assert detect_doc_date("plain.txt", ex) == (None, None)
