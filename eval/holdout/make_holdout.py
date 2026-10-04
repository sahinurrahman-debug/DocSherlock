"""Held-out evaluation corpus (a research lab, "Helix Biosciences"). Written AFTER the main system was tuned on the
Northwind demo corpus, with different domain vocabulary, so it measures generalisation rather than fit."""
from __future__ import annotations

import sys
from email.message import EmailMessage
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = HERE / "corpus"
OUT.mkdir(exist_ok=True)
sys.path.insert(0, str(HERE.parent.parent / "sample-documents"))
from generate import html_to_pdf  # noqa: E402


def main():
    html_to_pdf("""
    <h1>Laboratory Safety SOP</h1>
    <p class="meta">Version 2.0 &middot; Effective: January 10, 2023</p>
    <h2>Personal Protective Equipment</h2>
    <p>Lab coats must be worn at all times in the Wet Labs. Safety goggles are required when handling corrosive reagents.</p>
    <h2>Emergency Equipment</h2>
    <p>Eyewash stations must be tested weekly by the duty technician. The facility is staffed by 3 safety officers.</p>
    <h2>Spill Response</h2>
    <p>Chemical spills larger than 500 mL must be reported to the Safety Officer within 15 minutes of discovery.</p>
    """, OUT / "Lab_Safety_SOP_v2.pdf", "Laboratory Safety SOP")

    import docx
    d = docx.Document()
    d.add_heading("Safety Procedure Update", 0)
    d.add_paragraph("Issued: February 12, 2025. This memo revises the Laboratory Safety SOP.")
    d.add_heading("Emergency Equipment", 1)
    d.add_paragraph("With effect from March 1, 2025, eyewash stations must be tested every day.")
    d.add_heading("Spill Response", 1)
    d.add_paragraph("Spills larger than 250 mL must be reported to the Safety Officer within 30 minutes.")
    d.save(OUT / "Safety_Update_Memo_2025.docx")

    (OUT / "Grant_Award_Letter.txt").write_text(
        "GRANT AWARD LETTER\nDate: May 20, 2024\n\n"
        "AWARD\nThe Meridian Foundation awards Helix Biosciences a grant of $1,250,000 over 36 months.\n\n"
        "PAYMENT SCHEDULE\nThe first payment of $250,000 will be released on July 1, 2024. Remaining payments follow each January and July.\n", encoding="utf-8")

    (OUT / "Budget_FY2025.csv").write_text("Category,Amount ($)\nEquipment,420000\nPersonnel,610000\nTravel,35000\n", encoding="utf-8")

    m = EmailMessage()
    m["From"], m["To"], m["Date"], m["Subject"] = "Lena Ortiz <lena@helix.example>", "Board <board@helix.example>", "Tue, 15 Oct 2024 09:00:00 +0000", "Quarterly update"
    m.set_content("Hi all,\n\nGood news on funding: the Meridian Foundation grant totals $1.5 million over 36 months, which lets us start the sequencing project.\n\nBest,\nLena\n")
    (OUT / "Quarterly_Update_Email.eml").write_bytes(bytes(m))

    (OUT / "Staff_Directory.html").write_text(
        "<html><head><title>Helix staff</title></head><body><h1>Our team</h1>"
        "<p>Dr. Mara Linde is the Director of Research at Helix Biosciences. The lab has 24 researchers.</p></body></html>", encoding="utf-8")
    (OUT / "Annual_Review_2024.md").write_text(
        "# Annual Review 2024\n\nPublished December 2024.\n\n## People\n\nHelix employs 31 researchers across three teams.\n", encoding="utf-8")


QUESTIONS = [
    {"q": "How often must eyewash stations be tested?", "expect": "conflict", "values": ["weekly", "every day"], "docs": ["Lab_Safety_SOP_v2.pdf", "Safety_Update_Memo_2025.docx"]},
    {"q": "What is the spill reporting threshold?", "expect": "conflict", "values": ["500 mL", "250 mL"], "docs": ["Lab_Safety_SOP_v2.pdf", "Safety_Update_Memo_2025.docx"]},
    {"q": "Within how many minutes must a large spill be reported?", "expect": "conflict", "values": ["15 minutes", "30 minutes"], "docs": ["Lab_Safety_SOP_v2.pdf", "Safety_Update_Memo_2025.docx"]},
    {"q": "How much is the Meridian Foundation grant?", "expect": "conflict", "values": ["$1,250,000", "$1.5 million"], "docs": ["Grant_Award_Letter.txt", "Quarterly_Update_Email.eml"]},
    {"q": "How many researchers does Helix have?", "expect": "conflict", "values": ["24", "31"], "docs": ["Staff_Directory.html", "Annual_Review_2024.md"]},
    {"q": "How long is the grant period?", "expect": "answered", "contains": ["36 months"], "docs": ["Grant_Award_Letter.txt"]},
    {"q": "When is the first grant payment released?", "expect": "answered", "headline": "July 1, 2024", "docs": ["Grant_Award_Letter.txt"]},
    {"q": "Who is the Director of Research?", "expect": "answered", "contains": ["Mara Linde"], "docs": ["Staff_Directory.html"]},
    {"q": "What is the personnel budget?", "expect": "answered", "headline": "$610000", "docs": ["Budget_FY2025.csv"]},
    {"q": "What is the travel budget?", "expect": "answered", "headline": "$35000", "docs": ["Budget_FY2025.csv"]},
    {"q": "How many safety officers staff the facility?", "expect": "answered", "contains": ["3 safety officers"], "docs": ["Lab_Safety_SOP_v2.pdf"]},
    {"q": "Are lab coats required in the Wet Labs?", "expect": "answered", "contains": ["lab coats must be worn"], "docs": ["Lab_Safety_SOP_v2.pdf"]},
    {"q": "What is the airspeed velocity of an unladen swallow?", "expect": "insufficient"},
    {"q": "Who chairs the ethics committee?", "expect": "insufficient"},
    {"q": "What is the parking policy for visitors?", "expect": "insufficient"},
]

GROUND_TRUTH = [
    {"name": "eyewash frequency", "docs": ["Lab_Safety_SOP_v2.pdf", "Safety_Update_Memo_2025.docx"], "values": ["weekly", "every day"]},
    {"name": "spill volume", "docs": ["Lab_Safety_SOP_v2.pdf", "Safety_Update_Memo_2025.docx"], "values": ["500", "250"]},
    {"name": "spill reporting time", "docs": ["Lab_Safety_SOP_v2.pdf", "Safety_Update_Memo_2025.docx"], "values": ["15 minutes", "30 minutes"]},
    {"name": "grant amount", "docs": ["Grant_Award_Letter.txt", "Quarterly_Update_Email.eml"], "values": ["$1,250,000", "$1.5 million"]},
    {"name": "researcher headcount", "docs": ["Staff_Directory.html", "Annual_Review_2024.md"], "values": ["24", "31"]},
]

if __name__ == "__main__":
    import json
    main()
    (HERE / "questions.json").write_text(json.dumps(QUESTIONS, indent=1), encoding="utf-8")
    (HERE / "ground_truth_conflicts.json").write_text(json.dumps(GROUND_TRUTH, indent=1), encoding="utf-8")
    print("holdout corpus written to", OUT)
