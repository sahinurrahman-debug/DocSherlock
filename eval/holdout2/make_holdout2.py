"""Second held-out corpus (civil-engineering project). Written after the system was tuned on corpora 1 and 2 and run ONCE
with no further tuning; includes traps where a naive detector would raise false conflicts (actual-vs-spec values, workers-vs-visitors rules)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = HERE / "corpus"
OUT.mkdir(exist_ok=True)
sys.path.insert(0, str(HERE.parent.parent / "sample-documents"))
from generate import html_to_pdf  # noqa: E402


def main():
    html_to_pdf("""
    <h1>Riverside Bridge Project Charter</h1>
    <p class="meta">Issued: January 15, 2025</p>
    <h2>Budget and Schedule</h2>
    <p>The approved project budget is $12.4 million. The completion deadline is November 30, 2025.</p>
    <h2>Team</h2>
    <p>The project manager is Tomas Reyes. Maximum working hours are 10 hours per day for site crews.</p>
    """, OUT / "Project_Charter.pdf", "Riverside Bridge Project Charter")

    import docx
    d = docx.Document()
    d.add_heading("Change Order 7", 0)
    d.add_paragraph("Approved on August 3, 2025.")
    d.add_heading("Schedule and Cost Impact", 1)
    d.add_paragraph("Because of river-level delays the completion deadline is extended to February 28, 2026. The project budget is increased to $13.1 million.")
    d.save(OUT / "Change_Order_07.docx")

    (OUT / "Site_Inspection_Report.md").write_text(
        "# Site Inspection Report\n\nInspection date: July 22, 2025\n\n## Findings\n\n"
        "The concrete compressive strength achieved 42 MPa. The minimum specification is 40 MPa. Three safety violations were recorded during the visit.\n", encoding="utf-8")

    (OUT / "Weekly_Meeting_Notes.txt").write_text(
        "Weekly meeting notes - 2 September 2025\n\n"
        "Budget is still $12.4 million as far as I know. Tomas thinks we will finish by the end of January. The crane rental is a big cost line.\n", encoding="utf-8")

    (OUT / "Equipment_Rates.csv").write_text("Item,Daily rate ($)\nExcavator,420\nCrane,1200\n", encoding="utf-8")

    d = docx.Document()
    d.add_heading("Site Safety Plan", 0)
    d.add_paragraph("All workers must complete the safety induction before site access is granted.")
    d.add_paragraph("Hard hats are mandatory for all workers on site. The noise limit is 85 dB at the site boundary.")
    d.save(OUT / "Safety_Plan.docx")

    (OUT / "Community_Notice.html").write_text(
        "<html><head><title>Community notice</title></head><body><h1>Visiting the bridge works</h1>"
        "<p>Hard hats are optional for visitors in the public viewing area.</p></body></html>", encoding="utf-8")


QUESTIONS = [
    {"q": "What is the project budget?", "expect": "conflict", "values": ["$12.4 million", "$13.1 million"], "docs": ["Project_Charter.pdf", "Change_Order_07.docx"]},
    {"q": "When is the project due to be completed?", "expect": "conflict", "values": ["November 30, 2025", "February 28, 2026"], "docs": ["Project_Charter.pdf", "Change_Order_07.docx"]},
    {"q": "What was the concrete compressive strength?", "expect": "answered", "contains": ["42 MPa"], "docs": ["Site_Inspection_Report.md"]},
    {"q": "How many safety violations were recorded?", "expect": "answered", "contains": ["Three safety violations"], "docs": ["Site_Inspection_Report.md"]},
    {"q": "Who is the project manager?", "expect": "answered", "contains": ["Tomas Reyes"], "docs": ["Project_Charter.pdf"]},
    {"q": "What is the daily rate for an excavator?", "expect": "answered", "headline": "$420", "docs": ["Equipment_Rates.csv"]},
    {"q": "What is the noise limit at the site boundary?", "expect": "answered", "contains": ["85 dB"], "docs": ["Safety_Plan.docx"]},
    {"q": "Are hard hats mandatory for workers on site?", "expect": "answered", "contains": ["mandatory"], "docs": ["Safety_Plan.docx"]},
    {"q": "What are the maximum working hours per day?", "expect": "answered", "headline": "10 hours", "docs": ["Project_Charter.pdf"]},
    {"q": "What is the minimum concrete strength specification?", "expect": "answered", "contains": ["40 MPa"], "docs": ["Site_Inspection_Report.md"]},
    {"q": "What is the project's carbon footprint?", "expect": "insufficient"},
    {"q": "Who is the structural engineer?", "expect": "insufficient"},
    {"q": "What is the daily rate for a bulldozer?", "expect": "insufficient"},
]

GROUND_TRUTH = [
    {"name": "budget", "docs": ["Project_Charter.pdf", "Change_Order_07.docx"], "values": ["$12.4 million", "$13.1 million"]},
    {"name": "deadline", "docs": ["Project_Charter.pdf", "Change_Order_07.docx"], "values": ["November 30, 2025", "February 28, 2026"]},
]

if __name__ == "__main__":
    main()
    (HERE / "questions.json").write_text(json.dumps(QUESTIONS, indent=1), encoding="utf-8")
    (HERE / "ground_truth_conflicts.json").write_text(json.dumps(GROUND_TRUTH, indent=1), encoding="utf-8")
    print("holdout2 corpus written to", OUT)
